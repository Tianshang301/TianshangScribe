"""0.9.0 P3 tests: refined error codes, error_response refinements, dry-run plan."""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.utils.exceptions import CellCoordinatesException

from tianshang_scribe.mcp.errors import (
    DOCUMENTATION_URL,
    McpErrorCode,
    error_response,
)
from tianshang_scribe.mcp.schemas import ToolOptions
from tianshang_scribe.mcp.tools._dryrun import (
    build_content_plan,
    build_edit_plan,
    classify_engine_error,
    estimate_range_cells,
    validate_cell_ref,
    validate_range,
)
from tianshang_scribe.mcp.tools.create import create_office_document
from tianshang_scribe.mcp.tools.excel_edit import edit_excel_workbook
from tianshang_scribe.mcp.tools.ppt_create import create_presentation
from tianshang_scribe.mcp.tools.ppt_edit import edit_presentation


def _book(path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws.append(['k', 'v'])
    wb.save(path)
    return path


def _deck(path: str) -> str:
    res = create_presentation(path, slides=[{'title': 'T'}])
    assert res['success'] is True, res
    return path


class TestErrorRefinements:
    def test_new_codes_exist(self) -> None:
        assert McpErrorCode.EXCEL_INVALID_CELL_REF == 1007
        assert McpErrorCode.EXCEL_INVALID_RANGE == 1008
        assert McpErrorCode.PPT_INVALID_SLIDE_INDEX == 1009
        assert McpErrorCode.EXCEL_SHEET_NOT_FOUND == 1010

    def test_payload_backward_compatible_without_refinements(self) -> None:
        res = error_response(McpErrorCode.INVALID_PARAMETER, 'boom')
        assert set(res) == {'success', 'error_code', 'error_message', 'suggested_fix', 'retryable'}

    def test_field_and_documentation_url_included_when_given(self) -> None:
        res = error_response(
            McpErrorCode.EXCEL_INVALID_CELL_REF,
            '"AAAA1" is not an A1-style reference',
            field='operations[0].cell',
            documentation_url=DOCUMENTATION_URL,
        )
        assert res['field'] == 'operations[0].cell'
        assert res['documentation_url'].endswith('#error-handling')


class TestValidators:
    @pytest.mark.parametrize('ref', ['A1', 'XFD1048576', 'b2'])
    def test_valid_cell_refs(self, ref: str) -> None:
        ok, detail = validate_cell_ref(ref)
        assert ok is True, detail

    @pytest.mark.parametrize('ref', ['A0', '1A', 'AAAA1', '', '=SUM(A1)', 'XFE1'])
    def test_invalid_cell_refs(self, ref: str) -> None:
        ok, detail = validate_cell_ref(ref)
        assert ok is False
        assert detail

    @pytest.mark.parametrize('rng', ['A1:C10', 'B2'])
    def test_valid_ranges(self, rng: str) -> None:
        ok, detail = validate_range(rng)
        assert ok is True, detail

    @pytest.mark.parametrize('rng', ['A1:B', ':A5', 'A1:C10:D15', '2:5'])
    def test_invalid_ranges(self, rng: str) -> None:
        ok, detail = validate_range(rng)
        assert ok is False
        assert detail

    def test_estimate_range_cells(self) -> None:
        assert estimate_range_cells('A1:C10') == 30
        assert estimate_range_cells('B2') == 1
        assert estimate_range_cells('junk') == 0


class TestClassifyEngineError:
    def test_slide_index_indexerror(self) -> None:
        exc = IndexError('slide_index out of range: 99')
        assert classify_engine_error(exc) == (McpErrorCode.PPT_INVALID_SLIDE_INDEX, 'slide_index')

    def test_sheet_not_found(self) -> None:
        exc = ValueError('Sheet not found: Nope')
        assert classify_engine_error(exc) == (McpErrorCode.EXCEL_SHEET_NOT_FOUND, 'sheet_name')

    def test_invalid_row_range(self) -> None:
        exc = ValueError('Invalid row range: \'x\' (expected "2:5")')
        assert classify_engine_error(exc) == (McpErrorCode.EXCEL_INVALID_RANGE, 'range')

    def test_bad_cell_coordinates_exception(self) -> None:
        assert classify_engine_error(CellCoordinatesException('bad')) == (
            McpErrorCode.EXCEL_INVALID_CELL_REF,
            'cell',
        )

    def test_unmatched_returns_none(self) -> None:
        assert classify_engine_error(ValueError('password required')) is None


class TestBuildEditPlanExcel:
    def test_bad_cell_ref_reported_with_code_1007(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        ops = [{'action': 'write_cell', 'cell': 'AAAA1', 'value': 1}]
        plan = build_edit_plan(book, ops)
        assert plan['all_valid'] is False
        finding = plan['validations'][0]
        assert finding['error_code'] == McpErrorCode.EXCEL_INVALID_CELL_REF
        assert finding['field'] == 'cell'

    def test_unknown_sheet_reported_with_code_1010(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        ops = [{'action': 'write_cell', 'cell': 'A1', 'value': 1, 'sheet_name': 'Nope'}]
        plan = build_edit_plan(book, ops)
        assert plan['validations'][0]['error_code'] == McpErrorCode.EXCEL_SHEET_NOT_FOUND

    def test_group_rows_rejects_column_form(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        ops = [{'action': 'group_rows', 'range': 'B:D'}]
        plan = build_edit_plan(book, ops)
        assert plan['validations'][0]['error_code'] == McpErrorCode.EXCEL_INVALID_RANGE

    def test_add_sheet_makes_later_references_valid(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        ops = [
            {'action': 'add_sheet', 'sheet_name': 'Fresh'},
            {'action': 'write_cell', 'cell': 'A1', 'value': 1, 'sheet_name': 'Fresh'},
        ]
        plan = build_edit_plan(book, ops)
        assert plan['all_valid'] is True, plan['validations']

    def test_impact_estimate_counts_ranges_and_rows(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        ops = [
            {'action': 'write_cell', 'cell': 'A1', 'value': 1},
            {'action': 'group_rows', 'range': '2:5'},
            {'action': 'number_format', 'number_format': 'B2:D4=0.00'},
        ]
        plan = build_edit_plan(book, ops)
        assert plan['estimated_impacted_cells'] == 1 + 4 + 9


class TestBuildEditPlanPpt:
    def test_out_of_bounds_slide_reported_with_code_1009(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        deck = _deck(str(tmp_path / 'd.pptx'))
        ops = [{'action': 'add_notes', 'slide_index': 5, 'notes': 'x'}]
        plan = build_edit_plan(deck, ops)
        assert plan['validations'][0]['error_code'] == McpErrorCode.PPT_INVALID_SLIDE_INDEX

    def test_add_slide_extends_running_count(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        deck = _deck(str(tmp_path / 'd.pptx'))
        ops = [
            {'action': 'add_slide'},
            {'action': 'add_notes', 'slide_index': 1, 'notes': 'x'},
        ]
        plan = build_edit_plan(deck, ops)
        assert plan['all_valid'] is True, plan['validations']

    def test_unknown_action_reported_with_code_1006(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        deck = _deck(str(tmp_path / 'd.pptx'))
        plan = build_edit_plan(deck, [{'action': 'nonexistent_action'}])
        assert plan['validations'][0]['error_code'] == McpErrorCode.INVALID_PARAMETER


class TestDryRunWiring:
    def test_excel_dry_run_reports_findings(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        res = edit_excel_workbook(
            book,
            [{'action': 'write_cell', 'cell': 'ZZZZ99', 'value': 1}],
            options=ToolOptions(dry_run=True),
        )
        assert res['success'] is True
        data = res['data']
        assert data['all_valid'] is False
        assert data['validations'][0]['error_code'] == McpErrorCode.EXCEL_INVALID_CELL_REF

    def test_create_dry_run_reports_findings(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        out = tmp_path / 'new.xlsx'
        res = create_office_document(
            format='xlsx',
            content=[{'type': 'paragraph', 'text': 'hi', 'cell': 'nope'}],
            output_path=str(out),
            options=ToolOptions(dry_run=True),
        )
        assert res['success'] is True
        data = res['data']
        assert data['planned_content'] == 1  # legacy keys intact
        assert data['all_valid'] is False
        assert data['validations'][0]['field'] == 'cell'
        assert build_content_plan([])['all_valid'] is True


class TestLiveClassification:
    def test_live_slide_index_maps_to_1009_with_field(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        deck = _deck(str(tmp_path / 'd.pptx'))
        res = edit_presentation(deck, [{'action': 'add_notes', 'slide_index': 9, 'notes': 'x'}])
        assert res['success'] is False
        assert res['error_code'] == McpErrorCode.PPT_INVALID_SLIDE_INDEX
        assert res['field'] == 'operations[0].slide_index'
        assert res['documentation_url'].endswith('#error-handling')

    def test_live_unknown_action_keeps_1006_with_field(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        res = edit_excel_workbook(book, [{'action': 'bogus_action'}])
        assert res['success'] is False
        assert res['error_code'] == McpErrorCode.INVALID_PARAMETER
        assert res['field'] == 'operations[0].action'

    def test_unmatched_value_error_keeps_legacy_document_locked(self, tmp_path) -> None:  # type: ignore[no-untyped-def]
        book = _book(str(tmp_path / 'w.xlsx'))
        res = edit_excel_workbook(
            book,
            [{'action': 'set_page_setup', 'orientation': 'diagonal'}],
        )
        assert res['success'] is False
        assert res['error_code'] == McpErrorCode.DOCUMENT_LOCKED
